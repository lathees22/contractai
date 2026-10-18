import re
import io
from typing import List, Dict, Any, Tuple
from pypdf import PdfReader
from docx import Document

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts raw text from a PDF file using pypdf."""
    reader = PdfReader(io.BytesIO(file_bytes))
    full_text = []
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            full_text.append(f"--- [Page {idx + 1}] ---\n" + page_text.strip())
    return "\n\n".join(full_text)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text and table contents from a DOCX file."""
    doc = Document(io.BytesIO(file_bytes))
    full_text = []
    
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            full_text.append(t)
            
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                full_text.append(f"[Table Row]: {row_text}")
                
    return "\n\n".join(full_text)

def split_into_numbered_sections(text: str) -> List[Dict[str, str]]:
    """
    Splits contract text into numbered, structured sections.
    Recognizes patterns such as:
    - Section 1. Definitions
    - 1. PAYMENT TERMS
    - ARTICLE III - TERMINATION
    - 2.1 Confidential Information
    """
    lines = text.split("\n")
    sections: List[Dict[str, str]] = []
    
    current_sec_num = "Preamble"
    current_sec_title = "Agreement Header & Recitals"
    current_content: List[str] = []
    
    # Header patterns
    header_regexes = [
        # Match 'SECTION 1. TERMINATION' or 'ARTICLE 2: GOVERNING LAW'
        re.compile(r'^\s*(?:SECTION|ARTICLE|CLAUSE)\s+([0-9IVXLCDM]+)[\.\:\-\s]+(.*)', re.IGNORECASE),
        # Match '1. PAYMENT TERMS' or '2.0 TERM AND TERMINATION' or '2.1 Confidentiality'
        re.compile(r'^\s*([0-9]{1,2}(?:\.[0-9]{1,2})*)\.?\s+([A-Z][A-Za-z0-9\s,\-\/\(\)]+)\s*$'),
        # Match standalone capital headings like 'INDEMNIFICATION & LIABILITY'
        re.compile(r'^\s*([A-Z\s\&\,\-]{4,40})\s*$')
    ]
    
    auto_counter = 1
    
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            if current_content:
                current_content.append("")
            continue
            
        matched = False
        new_num = None
        new_title = None
        
        # Check if line matches section headers
        for idx, regex in enumerate(header_regexes):
            m = regex.match(line)
            if m:
                if idx == 0:
                    new_num = f"Section {m.group(1).strip()}"
                    new_title = m.group(2).strip() or "General Provisions"
                    matched = True
                    break
                elif idx == 1:
                    new_num = f"Section {m.group(1).strip()}"
                    new_title = m.group(2).strip()
                    matched = True
                    break
                elif idx == 2 and len(line) >= 4 and len(line) <= 40:
                    # Only treat as section if it looks like a standard contract title
                    known_titles = ["DEFINITIONS", "SERVICES", "PAYMENT", "TERM", "TERMINATION", 
                                    "CONFIDENTIALITY", "LIABILITY", "INDEMNITY", "INDEMNIFICATION",
                                    "WARRANTY", "WARRANTIES", "GOVERNING LAW", "DISPUTE RESOLUTION",
                                    "FORCE MAJEURE", "INTELLECTUAL PROPERTY", "MISCELLANEOUS",
                                    "GENERAL", "NON-COMPETE", "REPRESENTATIONS"]
                    if any(kt in line.upper() for kt in known_titles):
                        new_num = f"Section {auto_counter}"
                        auto_counter += 1
                        new_title = line.title()
                        matched = True
                        break

        if matched:
            # Save accumulated section if it has content
            content_str = "\n".join(current_content).strip()
            if content_str:
                sections.append({
                    "section_number": current_sec_num,
                    "section_title": current_sec_title,
                    "content": content_str
                })
            current_sec_num = new_num
            current_sec_title = new_title
            current_content = []
        else:
            current_content.append(raw_line)
            
    # Save the last section
    content_str = "\n".join(current_content).strip()
    if content_str:
        sections.append({
            "section_number": current_sec_num,
            "section_title": current_sec_title,
            "content": content_str
        })
        
    # If no sections matched (e.g. single block of text), chunk by double newlines or paragraphs
    if len(sections) <= 1 and len(text) > 300:
        paras = [p.strip() for p in text.split("\n\n") if p.strip()]
        if len(paras) > 1:
            sections = []
            for i, p in enumerate(paras):
                first_line = p.split("\n")[0][:45]
                sections.append({
                    "section_number": f"Section {i+1}",
                    "section_title": first_line.title() if len(first_line) > 5 else f"Clause Group {i+1}",
                    "content": p
                })
                
    return sections
